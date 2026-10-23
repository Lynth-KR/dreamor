#!/usr/bin/env python3
# SPDX-License-Identifier: Apache-2.0
"""run_doe.py results CSV -> analysis workbook.

  python3 make_report.py results/<run>/result.csv [report.xlsx]

The workbook goes beside the CSV as report.xlsx unless a path is given.

Every summary and effect in the workbook is a formula over the Raw sheet: add or fix rows in Raw
and it recalculates. Python writes only the design labels and the raw measurements as values.

Ratios are flat / dreamor (map / dreamor in the reference cell): above 1, dreamor is faster or
smaller. Effects are on ln(ratio), so they multiply. flat runs each set operation twice, hashed
(set_*) and sorted and merged (set_*_merge), dreamor once: both are set against dreamor's, and
set_*_best against whichever of flat's two is faster in each cell.

Sheets
  Summary   check result, the stage scenario and every response by N
  Scenario  the stage scenario by N x F and N x A, the reference cell with map, where the time
            goes, the stress run
  Design    factors and levels, arms, responses, the scenario's weights, fairness, machine
  Cells     the factorial cells x arms: geometric means, ratios, variances per response
  Effects   per response: level means, main effects (linear and quadratic for 3-level
            factors), two-factor interactions, standard errors from pure error, t
  Check     whether the arms handed out the same variables in the same groups, per model
  Charts    scenario ratio by N and F; residuals against run order (drift)
  Raw       measurements + derived columns (formulas)
"""

import csv
import itertools
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

FONT = "Arial"
BASE = "dreamor"
OTHER = "flat"  # the factorial's other arm
INAME = {
    "dreamor": "dreamor",
    "map": "std::map (hand-written)",
    "flat": "flat arrays (hand-optimized)",
}
# key, name, levels in design order (low ... high), what the high level exercises
FACTORS = [
    ("N_target", "rows N", ["100000", "1000000", "10000000"], "cache-resident -> DRAM-bound"),
    ("O", "K2 owner", ["shared", "owned"], "K2s that belong to one K1: many small K2 groups"),
    ("D", "density", ["dense", "sparse"], "10% of candidates kept, predicate evaluated, 10x the K1s"),
    ("S", "fan-out", ["uniform", "skewed"], "Pareto-distributed list sizes with the same mean"),
    (
        "A",
        "id order",
        ["aligned", "inner", "all"],
        "ids that do not follow the order rows come in: the K2/K3 lists and ids, then the K1 ids too",
    ),
    ("G", "group size", ["large", "small"], "5 K2s and 10 K3s per K1 instead of 50 and 100: many small groups"),
    ("F", "families", ["1", "3"], "three variable families on the same set"),
]
PHASES = [
    "build",
    "prep_contig",
    "prep_skip",
    "prep_lead",
    "g_contig",
    "g_skip",
    "g_lead",
    "slice",
    "coef",
    "cross",
    "neighbor",
    "set_select",
    "set_select_merge",
    "set_unite",
    "set_unite_merge",
    "set_intersect",
    "set_intersect_merge",
    "set_difference",
    "set_difference_merge",
    "set_member",
    "set_member_merge",
]
# flat's second way of each set operation -> the operation, which is dreamor's only way.
MERGED = {p: p.removesuffix("_merge") for p in PHASES if p.endswith("_merge")}
# The stage (examples/production_planning.hpp): the phases its model build is made of, once each.
SCENARIO = ["build", "prep_skip", "g_skip", "prep_lead", "slice", "coef", "cross", "neighbor"]
PHASE_DESC = {
    "build": "rows + container + F variable families",
    "prep_contig": "grouping by (K1, K2): the innermost column dropped, groups are runs",
    "prep_skip": "grouping by (K1, K3): the middle column dropped",
    "prep_lead": "grouping by (K2, K3): the leading column dropped",
    "g_contig": "one expression per (K1, K2) group, the table as the term",
    "g_skip": "one expression per (K1, K3) group, coef2(k2) * x looked up by columns",
    "g_lead": "one expression per (K2, K3) group, the table as the term",
    "slice": "a driver set of (K2, K3) keys, one expression per key from the (K2, K3) grouping",
    "coef": "one expression over every row: coef1(k1) * y, y the last family",
    "cross": "every row's x and the y of the same row, found by the row",
    "neighbor": "every row's neighbour (k1, k2, previous k3), looked up by its columns",
    "set_select": "the set projected on (K2, K3), first occurrences",
    "set_unite": "A united with B: A = the rows inA keeps, B = rows generated apart from moved K3 lists",
    "set_intersect": "A intersected with B",
    "set_difference": "A without B",
    "set_member": "a membership of B built, and the set filtered by it",
}
for p, op in MERGED.items():
    PHASE_DESC[p] = f"{op}, flat sorting the packed keys and merging (<algorithm>), against dreamor's {op}"
# key, name, Raw source column (None: worked out per cell, in Cells), unit, definition
RESPONSES = [("scenario", "stage scenario", "scenario_ms", "ms", "the stage's phases, summed: " + " + ".join(SCENARIO))]
for p in PHASES:
    RESPONSES.append((p, p, f"{p}_ms", "ms", PHASE_DESC[p]))
    if p in MERGED:
        op = MERGED[p]
        RESPONSES.append((f"{op}_best", f"{op} best", None, "ms", f"flat's faster of {op} and {p} in each cell"))
RESPONSES += [
    ("hit", "point lookup", "hit_ns", "ns", "one lookup of an existing key, random order"),
    ("miss", "absent lookup", "miss_ns", "ns", "one membership test of an absent key"),
    ("bytes", "memory", "bytes_row", "B/row", "resident growth over build + groupings, variables' heap excluded"),
]
# Every phase that hands out variables or rows has a check; build and the groupings do not. A merge
# phase's check must also equal its operation's.
CHECKS = ["groups"] + [f"chk_{p}" for p in PHASES if p != "build" and not p.startswith("prep_")]
TEXT_COLUMNS = {"impl", "O", "D", "S", "A", "G", "phase", "status"}

BLUE = Font(name=FONT, color="0000FF")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
RATIO = '0.00"x"'


def header(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


def note(ws, cell, text, merge_to, height=45):
    ws[cell] = text
    ws[cell].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(f"{cell}:{merge_to}")
    ws.row_dimensions[ws[cell].row].height = height


def finish_fonts(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.font is None or c.font.name != FONT:
                    f = c.font
                    c.font = Font(name=FONT, bold=f.bold, italic=f.italic, size=f.size, color=f.color)


def cell_id(levels):
    return "|".join(str(v) for v in levels)


def main():
    if len(sys.argv) not in (2, 3):
        sys.exit(__doc__)
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) == 3 else os.path.join(os.path.dirname(src), "report.xlsx")
    with open(src) as f:
        rows = list(csv.DictReader(f))
    if not rows:
        sys.exit(f"{src}: no rows")
    meta_path = os.path.splitext(src)[0] + "_meta.json"
    meta = json.load(open(meta_path)) if os.path.exists(meta_path) else {}
    ok_rows = [r for r in rows if r["status"] == "ok"]
    fact_rows = [r for r in ok_rows if r["phase"] == "factorial"]

    # The levels that ran, in design order; a factor with one level drops out of the effects.
    levels = {}
    for key, _, design, _ in FACTORS:
        present = {r[key] for r in fact_rows}
        levels[key] = [v for v in design if v in present] + sorted(present - set(design))
    varied = [f for f in FACTORS if len(levels[f[0]]) >= 2]
    run_cells = {tuple(r[f[0]] for f in FACTORS) for r in fact_rows}
    cells = [c for c in itertools.product(*[levels[f[0]] for f in FACTORS]) if c in run_cells]
    ref = meta.get("reference_cell") or dict(O="shared", D="dense", S="uniform", A="aligned", G="large", F="3")

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_scn = wb.create_sheet("Scenario")
    ws_design = wb.create_sheet("Design")
    ws_cell = wb.create_sheet("Cells")
    ws_eff = wb.create_sheet("Effects")
    ws_chk = wb.create_sheet("Check")
    ws_chart = wb.create_sheet("Charts")
    ws_raw = wb.create_sheet("Raw")

    # ------------------------------------------------------------------------------- Raw
    base_cols = list(rows[0].keys())
    header(ws_raw, 1, base_cols)
    col = {name: i + 1 for i, name in enumerate(base_cols)}
    for r, row in enumerate(rows, start=2):
        for name, v in row.items():
            if v in ("", None):
                continue  # an empty cell: the averages and MAXIFS skip it
            if name not in TEXT_COLUMNS:
                try:
                    v = float(v)
                    v = int(v) if v.is_integer() and abs(v) < 2**53 else v
                except ValueError:
                    pass
            ws_raw.cell(row=r, column=col[name], value=v)
    last = len(rows) + 1

    def rc(name):  # Raw column letter
        return L(col[name])

    ok = lambda r: f'{rc("status")}{r}="ok"'  # noqa: E731
    scenario_cols = [rc(p + "_ms") for p in SCENARIO]
    derived = [
        ("cell", lambda r: "=" + '&"|"&'.join(f"{rc(f[0])}{r}" for f in FACTORS)),
        ("scenario_ms", lambda r: f'=IF({ok(r)},SUM({",".join(c + str(r) for c in scenario_cols)}),"")'),
        ("bytes_row", lambda r: f'=IFERROR({rc("container_bytes")}{r}/{rc("N")}{r},"")'),
    ]
    for name, _ in derived:
        col[name] = len(col) + 1

    def ln_of(source, base):  # base: what dreamor's row reads instead (a merge phase's operation)
        def formula(r):
            v = f"{L(col[source])}{r}"
            if base != source:
                v = f'IF({rc("impl")}{r}="{BASE}",{L(col[base])}{r},{v})'
            return f'=IF({ok(r)},IFERROR(LN({v}),""),"")'

        return formula

    measured = [(key, source) for key, _, source, *_ in RESPONSES if source is not None]
    for key, source in measured:
        col[f"ln_{key}"] = len(col) + 1
        derived.append((f"ln_{key}", ln_of(source, f"{MERGED[key]}_ms" if key in MERGED else source)))
    for key, _ in measured:
        col[f"ln_{key}_sq"] = len(col) + 1
        derived.append((f"ln_{key}_sq", (lambda k: lambda r: f'=IFERROR({L(col[f"ln_{k}"])}{r}^2,"")')(key)))
    # Residual of the scenario against its cell and arm: drift over the run order shows here.
    col["resid_scenario"] = len(col) + 1
    lnc = rc("ln_scenario")
    derived.append(
        (
            "resid_scenario",
            lambda r: (
                f'=IF({rc("phase")}{r}="factorial",IFERROR({lnc}{r}-AVERAGEIFS(${lnc}$2:${lnc}${last},'
                f"${rc('cell')}$2:${rc('cell')}${last},{rc('cell')}{r},${rc('impl')}$2:${rc('impl')}${last},"
                f'{rc("impl")}{r},${rc("phase")}$2:${rc("phase")}${last},"factorial"),""),"")'
            ),
        )
    )
    start = len(base_cols) + 1
    header(ws_raw, 1, [d[0] for d in derived], col=start)
    for r in range(2, last + 1):
        for i, (_, f) in enumerate(derived):
            ws_raw.cell(row=r, column=start + i, value=f(r))
    ws_raw.freeze_panes = "B2"

    def rng(name):
        return f"Raw!${rc(name)}$2:${rc(name)}${last}"

    # ------------------------------------------------------------------------------- Cells
    ws = ws_cell
    ws["A1"] = (
        f"Cells - the {len(cells)} factorial cells (phase=factorial, status=ok). "
        "Geometric mean = EXP(mean of ln over repetitions); ratio = flat / dreamor."
    )
    ws["A1"].font = TITLE
    nf = len(FACTORS)
    hdr = [f[0] for f in FACTORS] + ["cell"]
    ccol = {name: i + 1 for i, name in enumerate(hdr)}

    # Contrasts: a 2-level factor is -1 / +1; a 3-level one has a linear (-1, 0, 1) and a quadratic
    # (1, -2, 1) part. Interactions take the first contrast of each factor.
    contrasts = []  # (code, factor key, label, values by level)
    for key, name, *_ in varied:
        lv = levels[key]
        if len(lv) == 2:
            contrasts.append((f"c_{key}", key, f"{key} {name}", {lv[0]: -1, lv[1]: 1}))
        else:
            contrasts.append((f"c_{key}l", key, f"{key} {name} (linear)", dict(zip(lv, [-1, 0, 1]))))
            contrasts.append((f"c_{key}q", key, f"{key} {name} (quadratic)", dict(zip(lv, [1, -2, 1]))))
    primary = {}
    for code, key, *_ in contrasts:
        primary.setdefault(key, code)
    pairs = list(itertools.combinations([f[0] for f in varied], 2))
    for code, *_ in contrasts:
        ccol[code] = len(ccol) + 1
        hdr.append(code)
    for a, b in pairs:
        ccol[f"x_{a}_{b}"] = len(ccol) + 1
        hdr.append(f"x_{a}_{b}")
    for c in (BASE, OTHER):
        ccol[f"n_{c}"] = len(ccol) + 1
        hdr.append(f"n_{c}")
    for key, *_ in RESPONSES:
        for c in (BASE, OTHER):
            for stat in ("m", "v"):
                ccol[f"{stat}_{key}_{c}"] = len(ccol) + 1
                hdr.append(f"{stat}_{key}_{c}")
    for key, rname, _, unit, _ in RESPONSES:
        for c in (BASE, OTHER):
            ccol[f"g_{key}_{c}"] = len(ccol) + 1
            hdr.append(f"{rname} {c} ({unit})")
        ccol[f"r_{key}"] = len(ccol) + 1
        hdr.append(f"{rname} flat / dreamor")
        ccol[f"d_{key}"] = len(ccol) + 1
        hdr.append(f"d_{key}")
        ccol[f"vd_{key}"] = len(ccol) + 1
        hdr.append(f"vd_{key}")
    header(ws, 2, hdr)
    R0 = 3
    for i, lv in enumerate(cells):
        r = R0 + i
        for j, v in enumerate(lv):
            ws.cell(row=r, column=j + 1, value=int(v) if FACTORS[j][0] in ("N_target", "F") else v)
        ws.cell(row=r, column=ccol["cell"], value="=" + '&"|"&'.join(f"{L(j + 1)}{r}" for j in range(nf)))
        for code, key, _, values in contrasts:
            j = [f[0] for f in FACTORS].index(key)
            ws.cell(row=r, column=ccol[code], value=values[lv[j]])
        for a, b in pairs:
            ws.cell(row=r, column=ccol[f"x_{a}_{b}"], value=f"={L(ccol[primary[a]])}{r}*{L(ccol[primary[b]])}{r}")
        crit = lambda c, r=r: (  # noqa: E731
            f'{rng("cell")},${L(ccol["cell"])}{r},{rng("impl")},"{c}",{rng("phase")},"factorial",{rng("status")},"ok"'
        )
        for c in (BASE, OTHER):
            ws.cell(row=r, column=ccol[f"n_{c}"], value=f"=COUNTIFS({crit(c)})")
        for key, _, source, *_ in RESPONSES:
            for c in (BASE, OTHER):
                n = f"{L(ccol['n_' + c])}{r}"
                m = f"{L(ccol[f'm_{key}_{c}'])}{r}"
                if source is None:  # a set operation's best: flat's faster way in this cell, dreamor's one way
                    op = key.removesuffix("_best")
                    mh, vh = (f"{L(ccol[f'{s}_{op}_{c}'])}{r}" for s in ("m", "v"))
                    mm, vm = (f"{L(ccol[f'{s}_{op}_merge_{c}'])}{r}" for s in ("m", "v"))
                    mean, var = f"={mh}", f"={vh}"
                    if c == OTHER:
                        mean = f'=IF(COUNT({mh},{mm})=0,"",MIN({mh},{mm}))'
                        var = f"=IF({m}={mh},{vh},{vm})"
                else:
                    mean = f'=IFERROR(AVERAGEIFS({rng("ln_" + key)},{crit(c)}),"")'
                    var = f'=IFERROR(IF({n}>1,(SUMIFS({rng("ln_" + key + "_sq")},{crit(c)})-{n}*{m}^2)/({n}-1),0),"")'
                ws.cell(row=r, column=ccol[f"m_{key}_{c}"], value=mean)
                ws.cell(row=r, column=ccol[f"v_{key}_{c}"], value=var)
                ws.cell(row=r, column=ccol[f"g_{key}_{c}"], value=f'=IFERROR(EXP({m}),"")').number_format = "#,##0.00"
            mo, mb = f"{L(ccol[f'm_{key}_{OTHER}'])}{r}", f"{L(ccol[f'm_{key}_{BASE}'])}{r}"
            vo, vb = f"{L(ccol[f'v_{key}_{OTHER}'])}{r}", f"{L(ccol[f'v_{key}_{BASE}'])}{r}"
            no, nb = f"{L(ccol['n_' + OTHER])}{r}", f"{L(ccol['n_' + BASE])}{r}"
            ws.cell(row=r, column=ccol[f"d_{key}"], value=f'=IFERROR({mo}-{mb},"")')
            ws.cell(
                row=r, column=ccol[f"r_{key}"], value=f'=IFERROR(EXP({L(ccol[f"d_{key}"])}{r}),"")'
            ).number_format = RATIO
            ws.cell(row=r, column=ccol[f"vd_{key}"], value=f'=IFERROR({vo}/{no}+{vb}/{nb},"")')
    R1 = R0 + len(cells) - 1
    for n in [n for n in ccol if n.startswith(("m_", "v_", "d_", "vd_", "c_", "x_"))]:
        ws.column_dimensions[L(ccol[n])].hidden = True
    ws.column_dimensions[L(ccol["cell"])].width = 44
    ws.freeze_panes = f"{L(ccol['cell'] + 1)}3"

    def cr(name):  # range in Cells
        return f"Cells!${L(ccol[name])}${R0}:${L(ccol[name])}${R1}"

    # ------------------------------------------------------------------------------- Effects
    ws = ws_eff
    ws["A1"] = "Effects - on ln(flat / dreamor), per response"
    ws["A1"].font = TITLE
    note(
        ws,
        "A2",
        "Level means: the geometric-mean ratio over the cells at that level. Effect of a contrast c = "
        "SUMPRODUCT(c, d) / SUM(c > 0): for a 2-level factor and a linear contrast, mean at the high level - "
        "mean at the low level; for a quadratic contrast, mean of the two ends - mean of the middle (curvature); "
        "for an interaction, how much one factor's effect changes with the other. 'Multiplier' = EXP(effect): "
        "how many times the ratio changes. Standard errors come from the variance between repetitions within "
        "cells (pure error); |t| above about 2 is larger than the noise.",
        "H2",
        height=75,
    )
    terms = [(label, code) for code, _, label, _ in contrasts] + [(f"{a} x {b}", f"x_{a}_{b}") for a, b in pairs]
    row = 4
    lvl_ref = {}  # (response, factor, level) -> cell with the level's geometric-mean ratio
    for key, rname, _, unit, _ in RESPONSES:
        ws.cell(row=row, column=1, value=f"{rname} ({unit}) - flat / dreamor").font = BOLD
        row += 1
        d, vd = cr(f"d_{key}"), cr(f"vd_{key}")
        header(ws, row, ["level means", "level", "ratio", "cells"])
        row += 1
        ws.cell(row=row, column=1, value="all cells")
        ws.cell(row=row, column=3, value=f'=IFERROR(EXP(AVERAGE({d})),"")').number_format = RATIO
        ws.cell(row=row, column=4, value=f"=COUNT({d})")
        lvl_ref[(key, None, None)] = f"Effects!$C${row}"
        row += 1
        for fkey, fname, *_ in varied:
            for lv in levels[fkey]:
                fcol = cr(fkey)
                crit_v = int(lv) if fkey in ("N_target", "F") else f'"{lv}"'
                ws.cell(row=row, column=1, value=f"{fkey} {fname}")
                ws.cell(row=row, column=2, value=lv)
                ws.cell(
                    row=row, column=3, value=f'=IFERROR(EXP(AVERAGEIFS({d},{fcol},{crit_v})),"")'
                ).number_format = RATIO
                ws.cell(row=row, column=4, value=f'=COUNTIFS({fcol},{crit_v},{d},"<>")')
                lvl_ref[(key, fkey, lv)] = f"Effects!$C${row}"
                row += 1
        header(ws, row, ["term", "effect (ln)", "multiplier", "std. error", "t"])
        row += 1
        for name, code in terms:
            c = cr(code)
            ws.cell(row=row, column=1, value=name)
            ws.cell(
                row=row, column=2, value=f'=IFERROR(SUMPRODUCT({c},{d})/SUMIF({c},">0"),"")'
            ).number_format = "0.000"
            ws.cell(row=row, column=3, value=f'=IFERROR(EXP(B{row}),"")').number_format = RATIO
            ws.cell(
                row=row, column=4, value=f'=IFERROR(SQRT(SUMPRODUCT({c}^2,{vd}))/SUMIF({c},">0"),"")'
            ).number_format = "0.000"
            ws.cell(row=row, column=5, value=f'=IFERROR(B{row}/D{row},"")').number_format = "0.0"
            row += 1
        row += 1
    ws.column_dimensions["A"].width = 34
    for c in "BCDE":
        ws.column_dimensions[c].width = 13

    # ------------------------------------------------------------------------------- Check
    ws = ws_chk
    ws["A1"] = "Check - did every arm hand out the same variables in the same groups, per model (cell, repetition)?"
    ws["A1"].font = TITLE
    note(
        ws,
        "A2",
        "groups = non-empty (K1, K3) groups. chk_<phase> = sum over the phase's expressions of (the terms' "
        "bound x coefficient, summed) x a hash of the phase and the group's key. Bounds are tags from each row's "
        "key and family, keyed by serials - not addresses, not ids - so equal values across arms mean the same "
        "variables, coefficients and groups. A phase an arm "
        "does not run is empty and not compared. A set operation's chk is order-free (its rows' hashes summed), "
        "and flat's chk_<operation>_merge must equal the operation's chk.",
        "M2",
        height=60,
    )
    hdr = ["cell", "rep"] + [x for c in CHECKS for x in (f"{c} max", f"{c} min")] + ["runs", "result"]
    header(ws, 4, hdr)
    keys = sorted({(cell_id(r[f[0]] for f in FACTORS), int(r["rep"])) for r in ok_rows})
    row = 5
    for cid, rep in keys:
        crit = f'{rng("cell")},$A{row},{rng("rep")},$B{row},{rng("status")},"ok"'
        ws.cell(row=row, column=1, value=cid)
        ws.cell(row=row, column=2, value=rep)
        conds = []
        for j, name in enumerate(CHECKS):
            cmax, cmin = 3 + 2 * j, 4 + 2 * j
            ws.cell(row=row, column=cmax, value=f"=_xlfn.MAXIFS({rng(name)},{crit})").number_format = "0"
            ws.cell(row=row, column=cmin, value=f"=_xlfn.MINIFS({rng(name)},{crit})").number_format = "0"
            same = f"{L(cmax)}{row}={L(cmin)}{row}"
            if (phase := name.removeprefix("chk_")) in MERGED:  # flat's alone: equal to its operation's, if it ran
                op_max = L(3 + 2 * CHECKS.index(f"chk_{MERGED[phase]}"))
                same = f'OR(COUNTIFS({rng(name)},"<>",{crit})=0,AND({same},{L(cmax)}{row}={op_max}{row}))'
            conds.append(same)
        ws.cell(row=row, column=len(hdr) - 1, value=f"=COUNTIFS({crit})")
        ws.cell(row=row, column=len(hdr), value=f'=IF(AND({",".join(conds)}),"match","MISMATCH")')
        row += 1
    chk_last = max(row - 1, 5)
    chk_result = L(len(hdr))
    ws.column_dimensions["A"].width = 44
    for j in range(3, len(hdr) - 1):
        ws.column_dimensions[L(j)].width = 22

    # ------------------------------------------------------------------------------- Scenario
    ws = ws_scn
    ws["A1"] = "Scenario - the stage's model build (examples/production_planning.hpp)"
    ws["A1"].font = TITLE
    note(
        ws,
        "A2",
        "The scenario sums the phases the example's model build is made of: " + " + ".join(SCENARIO) + ". "
        "Ratios are geometric means of flat / dreamor over the factorial cells in each group (above 1: dreamor "
        "is faster). The reference cell is the example's own shape: "
        + ", ".join(f"{k}={v}" for k, v in ref.items())
        + ".",
        "J2",
    )
    row = 4
    d_scn = cr("d_scenario")
    n_levels = levels["N_target"]

    def table(title, fkey):
        nonlocal row
        ws.cell(row=row, column=1, value=title).font = BOLD
        row += 1
        lv = levels[fkey]
        header(ws, row, ["N"] + [f"{fkey}={v}" for v in lv] + ["all"])
        row += 1
        for n in n_levels:
            ws.cell(row=row, column=1, value=int(n)).number_format = "#,##0"
            for j, v in enumerate(lv):
                crit_v = int(v) if fkey in ("N_target", "F") else f'"{v}"'
                ws.cell(
                    row=row,
                    column=2 + j,
                    value=f'=IFERROR(EXP(AVERAGEIFS({d_scn},{cr("N_target")},$A{row},{cr(fkey)},{crit_v})),"")',
                ).number_format = RATIO
            ws.cell(
                row=row, column=2 + len(lv), value=f'=IFERROR(EXP(AVERAGEIFS({d_scn},{cr("N_target")},$A{row})),"")'
            ).number_format = RATIO
            row += 1
        row += 1

    scn_table_row = row + 2
    table("Scenario ratio flat / dreamor by N and F (families on one set)", "F")
    table("Scenario ratio flat / dreamor by N and A (id order)", "A")
    table("Scenario ratio flat / dreamor by N and D (density)", "D")

    # Reference cell: all three arms, by N.
    ws.cell(row=row, column=1, value="Reference cell: scenario ms (geometric mean over repetitions)").font = BOLD
    row += 1
    header(ws, row, ["N", "dreamor", "flat", "map", "flat / dreamor", "map / dreamor"])
    row += 1
    ref_cell = lambda n: cell_id([n] + [ref[f[0]] for f in FACTORS[1:]])  # noqa: E731
    for n in n_levels:
        ws.cell(row=row, column=1, value=int(n)).number_format = "#,##0"
        for j, (c, phase) in enumerate([(BASE, "factorial"), (OTHER, "factorial"), ("map", "reference")]):
            crit = f'{rng("cell")},"{ref_cell(n)}",{rng("impl")},"{c}",{rng("phase")},"{phase}",{rng("status")},"ok"'
            ws.cell(
                row=row, column=2 + j, value=f'=IFERROR(EXP(AVERAGEIFS({rng("ln_scenario")},{crit})),"")'
            ).number_format = "#,##0.0"
        ws.cell(row=row, column=5, value=f'=IFERROR(C{row}/B{row},"")').number_format = RATIO
        ws.cell(row=row, column=6, value=f'=IFERROR(D{row}/B{row},"")').number_format = RATIO
        row += 1
    row += 1

    # Where the time goes, reference cell, per phase: wall and system time.
    ws.cell(
        row=row, column=1, value="Reference cell: where the time goes (mean ms; sys = kernel, mostly page faults)"
    ).font = BOLD
    row += 1
    hdr = ["N", "arm"] + [x for p in PHASES for x in (p, f"{p} sys")]
    header(ws, row, hdr)
    row += 1
    for n in n_levels:
        for c in (BASE, OTHER):
            crit = f'{rng("cell")},"{ref_cell(n)}",{rng("impl")},"{c}",{rng("phase")},"factorial",{rng("status")},"ok"'
            ws.cell(row=row, column=1, value=int(n)).number_format = "#,##0"
            ws.cell(row=row, column=2, value=c)
            for j, p in enumerate(PHASES):
                for k, suffix in enumerate(["_ms", "_sys_ms"]):
                    ws.cell(
                        row=row, column=3 + 2 * j + k, value=f'=IFERROR(AVERAGEIFS({rng(p + suffix)},{crit}),"")'
                    ).number_format = "#,##0.0"
            row += 1
    row += 1

    # Stress.
    ws.cell(
        row=row, column=1, value="Stress: the reference cell beyond the design's N (phase=stress, not analysed)"
    ).font = BOLD
    row += 1
    header(ws, row, ["N", "dreamor ms", "flat ms", "flat / dreamor"])
    row += 1
    if meta.get("stress_n"):
        n = str(meta["stress_n"])
        ws.cell(row=row, column=1, value=int(n)).number_format = "#,##0"
        for j, c in enumerate((BASE, OTHER)):
            crit = f'{rng("cell")},"{ref_cell(n)}",{rng("impl")},"{c}",{rng("phase")},"stress",{rng("status")},"ok"'
            ws.cell(
                row=row, column=2 + j, value=f'=IFERROR(EXP(AVERAGEIFS({rng("ln_scenario")},{crit})),"")'
            ).number_format = "#,##0.0"
        ws.cell(row=row, column=4, value=f'=IFERROR(C{row}/B{row},"")').number_format = RATIO
    else:
        ws.cell(row=row, column=1, value="(not run: run_doe.py --stress)")
    ws.column_dimensions["A"].width = 14
    for j in range(2, 26):
        ws.column_dimensions[L(j)].width = 14

    # ------------------------------------------------------------------------------- Design
    ws = ws_design
    ws["A1"] = "Design"
    ws["A1"].font = TITLE
    r = 3
    ws.cell(row=r, column=1, value="Model").font = BOLD
    r += 1
    for line in [
        "A 3-index set (K1, K2, K3) read as the lot-sizing example: K1 a machine, K2 an item, K3 a period. "
        "Rows = {(k1,k2,k3): k2 in C2(k1), k3 in C3(k1), keep}. One variable per row and family.",
        "K1s are added until there are N rows. Entities have a serial (their place in the data) and an id (the "
        "integer flat packs); A permutes ids only, so every level of A has the same rows and checks.",
        "Same repetition number = same seed = same model: the arms are measured on the same rows, pairwise; "
        "a repetition is a block.",
    ]:
        ws.cell(row=r, column=1, value=line)
        r += 1
    r += 1
    header(ws, r, ["factor", "name", "levels (low ... high)", "in this run", "what the higher levels exercise"])
    r += 1
    for key, name, design, why in FACTORS:
        for i, v in enumerate([key, name, " / ".join(design), " / ".join(levels[key]), why]):
            ws.cell(row=r, column=1 + i, value=v)
        r += 1
    r += 1
    header(ws, r, ["arm", "implementation", "groupings and sums"])
    r += 1
    for c, impl, path in [
        (
            "dreamor",
            "makeSet(...).where(keep).collect(), F VarTables on the set",
            "groupBy per grouping; sumOf(env, rows, vX) or a column lambda; slice() for the driver set",
        ),
        (
            "flat",
            "rows (ids) and F variable vectors; std::unordered_map on the packed ids (identity hash)",
            "(K1, K2) as runs; other groupings compacted by an unordered_map on the packed key, CSR by counting "
            "sort; direct indexing. Set operations twice: std::unordered_set, then sorted keys merged",
        ),
        (
            "map",
            "one std::map<tuple<E1*,E2*,E3*>, IloNumVar> per family, filled in one loop "
            "(examples/production_planning_plain.hpp); reference cell only",
            "loops over the data with keep() and a map lookup per variable; owner lists for the driver set",
        ),
    ]:
        for i, v in enumerate([INAME[c], impl, path]):
            ws.cell(row=r, column=1 + i, value=v)
        r += 1
    r += 1
    header(ws, r, ["response", "unit", "definition", "in the scenario"])
    r += 1
    for key, rname, _, unit, desc in RESPONSES:
        for i, v in enumerate([rname, unit, desc, "yes" if key in SCENARIO else ""]):
            ws.cell(row=r, column=1 + i, value=v)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Fairness").font = BOLD
    r += 1
    for line in [
        "flat is what a skilled hand-coder writes with the standard library, and keeps what they know: dense ids, rows in "
        "loop order, an identity hash on packed ids, (K1, K2) groups as runs. What that knowledge leans on - ids "
        "lining up with the order rows come in - is factor A. Its containers stay the standard library's, set "
        "operations included: each runs twice, the way that comes to hand (std::unordered_set) and the <algorithm> "
        "way (sorted keys merged, the result in key order rather than dreamor's first-appearance order), and "
        "set_*_best takes the faster per cell. A hand-written set library would be a second dreamor, not "
        "hand-written model code.",
        "No solver: bench/stub makes each variable a 16-byte heap object and each expression a vector of "
        "(variable, coefficient). Every arm consumes every expression the same way. A real solver's cost moves "
        "with its insides and the machine, not with the containers: bench/solver/ looks at it after the run.",
        "Phases run in a fixed order in a fresh process: every measurement is a cold one, as a model build is.",
        "Runs in a seeded random order; pinned to one CPU of the fastest core class unless virtualized "
        "(the Summary sheet names the machine).",
        "Memory is resident growth over build + groupings, minus the variables' own heap (measured at start-up).",
    ]:
        ws.cell(row=r, column=1, value=line)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Machine (recorded by run_doe.py)").font = BOLD
    r += 1
    for k in [
        "design",
        "impls",
        "reps",
        "seed",
        "cpu",
        "pinned_cpus",
        "cpu_choice",
        "governor",
        "boost",
        "smt_active",
        "thp",
        "llc_bytes",
        "glibc_tunables",
        "mem_total_bytes",
        "kernel",
        "compiler",
        "python",
        "commit",
        "dirty",
        "started",
    ]:
        ws.cell(row=r, column=1, value=k)
        v = meta.get(k, "")
        if isinstance(v, (list, dict)):
            v = json.dumps(v)
        c = ws.cell(row=r, column=2, value=v)
        if k in ("mem_total_bytes", "llc_bytes"):
            c.number_format = "#,##0"
        r += 1
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 24
    ws.column_dimensions["E"].width = 70

    # ------------------------------------------------------------------------------- Summary
    ws = ws_sum
    ws["A1"] = "Container benchmark - dreamor vs flat arrays (hand-optimized)"
    ws["A1"].font = TITLE
    ws["A3"] = "successful runs"
    ws["B3"] = f'=COUNTIFS({rng("status")},"ok")'
    ws["A4"] = "failed or timed out"
    ws["B4"] = f'=COUNTIFS({rng("status")},"<>ok")'
    ws["A5"] = "check mismatches"
    ws["B5"] = f'=COUNTIFS(Check!${chk_result}$5:${chk_result}${chk_last},"MISMATCH")'
    ws["C5"] = "must be 0"
    ws["A6"] = "factorial cells"
    ws["B6"] = len(cells)
    choice = meta.get("cpu_choice") or {}
    pinned = meta.get("pinned_cpus")
    ws["A7"] = "machine"
    ws["B7"] = (
        f"{meta.get('cpu', 'unknown CPU')} - {choice.get('virtualized') or 'bare metal'}, "
        f"{'pinned to CPU ' + ','.join(map(str, pinned)) if pinned else 'not pinned'}"
    )
    r = 8
    ws.cell(
        row=r,
        column=1,
        value="Geometric-mean ratio flat / dreamor over the cells at each N (above 1: dreamor is that many times "
        "faster or smaller). Effects sheet: what moves it.",
    ).font = BOLD
    r += 1
    header(ws, r, ["response", "unit"] + [f"N={int(n):,}" for n in n_levels] + ["all cells", "min cell", "max cell"])
    r += 1
    for key, rname, _, unit, _ in RESPONSES:
        name_cell = ws.cell(row=r, column=1, value=rname)
        if key == "scenario":
            name_cell.font = BOLD
        ws.cell(row=r, column=2, value=unit)
        for j, n in enumerate(n_levels):
            ref_cell_addr = lvl_ref.get((key, "N_target", n))
            ws.cell(
                row=r,
                column=3 + j,
                value=f"={ref_cell_addr}"
                if ref_cell_addr
                else f'=IFERROR(EXP(AVERAGEIFS({cr("d_" + key)},{cr("N_target")},{int(n)})),"")',
            ).number_format = RATIO
        k = 3 + len(n_levels)
        ws.cell(row=r, column=k, value=f"={lvl_ref[(key, None, None)]}").number_format = RATIO
        ws.cell(row=r, column=k + 1, value=f'=IFERROR(MIN({cr("r_" + key)}),"")').number_format = RATIO
        ws.cell(row=r, column=k + 2, value=f'=IFERROR(MAX({cr("r_" + key)}),"")').number_format = RATIO
        r += 1
    r += 1
    ws.cell(
        row=r,
        column=1,
        value="Scenario sheet: by N x F / A / D, the reference cell with map, where the time goes.",
    )
    ws.column_dimensions["A"].width = 30
    for c in "BCDEFGHI":
        ws.column_dimensions[c].width = 15

    # ------------------------------------------------------------------------------- Charts
    ws = ws_chart
    ws["A1"] = "Charts (values from the Scenario sheet and Raw)"
    ws["A1"].font = TITLE
    ch = BarChart()
    ch.title = "Stage scenario: flat / dreamor by N and F"
    ch.y_axis.title = "ratio"
    ch.height, ch.width = 12, 20
    hdr_row = scn_table_row - 1
    nlv = len(levels["F"])
    ch.add_data(
        Reference(ws_scn, min_col=2, max_col=1 + nlv, min_row=hdr_row, max_row=hdr_row + len(n_levels)),
        titles_from_data=True,
    )
    ch.set_categories(Reference(ws_scn, min_col=1, min_row=hdr_row + 1, max_row=hdr_row + len(n_levels)))
    ws.add_chart(ch, "A3")
    sc = ScatterChart()
    sc.title = "Drift: scenario residual (ln) against run order"
    sc.x_axis.title = "run order"
    sc.y_axis.title = "ln residual"
    sc.style = 13
    sc.height, sc.width = 12, 20
    s = Series(
        Reference(ws_raw, min_col=col["resid_scenario"], min_row=2, max_row=last),
        Reference(ws_raw, min_col=col["order"], min_row=2, max_row=last),
        title="factorial runs",
    )
    s.marker.symbol = "circle"
    s.marker.size = 3
    s.graphicalProperties.line.noFill = True
    sc.series.append(s)
    ws.add_chart(sc, "M3")

    finish_fonts(wb)
    out_dir = os.path.dirname(out)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(out)
    print(f"saved: {out}  (Raw {len(rows)} rows, {len(cells)} factorial cells)")


if __name__ == "__main__":
    main()
